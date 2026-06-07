"""
Add Denver County, Arapahoe County, and Douglas County to Resources sheet.
Update county dropdown to include all 9 counties.
"""

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

DST = "/home/user/Appraisal-Template/Analysis_Workbook_OPTIMIZED.xlsm"
wb = load_workbook(DST, keep_vba=True)
ws = wb["Resources"]

URL_FONT = Font(color="0563C1", underline="single", size=10)

def prio(r):
    return (
        f'=IF(B{r}<>$D$3,"◌ Other",'
        f'IF(C{r}="REGIONAL / ALL MUNICIPALITIES","Regional",'
        f'IF(OR('
        f'ISNUMBER(SEARCH(LEFT(\'Property Info\'!$B$17,MAX(1,FIND(",",\'Property Info\'!$B$17&",")-1)),C{r})),'
        f'ISNUMBER(SEARCH(\'Property Info\'!$B$18,C{r}))'
        f')),"★ Active","–")))'
    )

def append_row(county, municipality, category, name, url, description):
    r = ws.max_row + 1
    ws.cell(r, 1, prio(r))
    ws.cell(r, 2, county)
    ws.cell(r, 3, municipality)
    ws.cell(r, 4, category)
    ws.cell(r, 5, name)
    uc = ws.cell(r, 6, url)
    uc.hyperlink = url
    uc.font = URL_FONT
    ws.cell(r, 7, description)
    ws.row_dimensions[r].height = 15

# ── DENVER COUNTY ─────────────────────────────────────────────────────────────
DC  = "Denver County"
DCM = "Denver (City & County)"

denver_rows = [
    # Property Assessment
    (DC, DCM, "Property Assessment", "Denver Property Search", "https://www.denvergov.org/Property", "Property assessment and tax data lookup by address, parcel ID, or schedule number"),
    (DC, DCM, "Property Assessment", "Assessor's Office", "https://www.denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Department-of-Finance/Our-Divisions/Assessors-Office", "Official Denver Assessor — property valuation, appeals, and FAQs"),
    (DC, DCM, "Property Assessment", "Assessor Maps (GIS)", "https://www.denvergov.org/maps/map/assessormaps", "Interactive GIS maps showing assessment parcel data and property info"),
    (DC, DCM, "Property Assessment", "Commercial / Business Personal Property", "https://denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Department-of-Finance/Our-Divisions/Assessors-Office/Business-Personal-Property", "Business personal property valuation and assessment"),
    (DC, DCM, "Property Assessment", "Assessment FAQ & Appeals", "https://www.denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Department-of-Finance/Our-Divisions/Assessors-Office/Assessment-FAQ", "NOV appeal process, deadlines, and assessment FAQ"),
    (DC, DCM, "Property Assessment", "DOLA — Denver City & County Assessor", "https://dpt.colorado.gov/locality/denver-city-and-county-assessor", "State DPT profile page for Denver assessor"),
    # Property Tax / Treasurer
    (DC, DCM, "Property Tax / Treasurer", "Denver Property Taxes", "https://denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Department-of-Finance/Our-Divisions/Treasury/Property-Taxes", "Tax payment info, deadlines, and mill levy data"),
    # Clerk & Recorder
    (DC, DCM, "Clerk & Recorder / Official Records", "Denver Clerk & Recorder Records Search", "https://denver.co.ds.search.govos.com/", "Quick search for Denver County recorded documents"),
    (DC, DCM, "Clerk & Recorder / Official Records", "Clerk & Recorder Main", "https://www.denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Denver-Clerk-and-Recorder", "Denver Clerk and Recorder — recording and records services"),
    (DC, DCM, "Clerk & Recorder / Official Records", "Document Recording", "https://www.denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Denver-Clerk-and-Recorder/Recording-Division/Denver-Recording/document-recording", "Information on recording property documents"),
    # GIS / Mapping
    (DC, DCM, "GIS / Mapping", "Denver Maps Portal", "https://www.denvergov.org/maps", "Comprehensive GIS mapping portal with multiple interactive layers"),
    (DC, DCM, "GIS / Mapping", "Denver Zoning Map", "https://www.denvergov.org/Maps/map/zoning", "Interactive zoning map showing zone districts and designations"),
    (DC, DCM, "GIS / Mapping", "Denver Floodplain Map", "https://www.denvergov.org/maps/map/floodplain", "FEMA floodplain mapping and flood hazard areas"),
    (DC, DCM, "GIS / Mapping", "Denver Open Data Catalog", "https://denvergov.org/opendata", "Open geospatial data downloads — parcels, zoning, and property data"),
    (DC, DCM, "GIS / Mapping", "Denver Geospatial Hub (ArcGIS)", "https://opendata-geospatialdenver.hub.arcgis.com/", "Parcel and property datasets available for download"),
    # Planning & Zoning
    (DC, DCM, "Planning & Zoning", "Denver Zoning Code", "https://www.denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Community-Planning-and-Development/Denver-Zoning-Code", "Official Denver Zoning Code with zone district descriptions"),
    (DC, DCM, "Planning & Zoning", "Denver Zoning Code — Zone Descriptions", "https://denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Community-Planning-and-Development/Denver-Zoning-Code/Zone-Descriptions", "Detailed descriptions of all zone districts"),
    (DC, DCM, "Planning & Zoning", "Comprehensive Plan 2040 / Blueprint Denver", "https://denvergov.org/Government/Citywide-Programs-and-Initiatives/Comprehensive-Plan-2040", "20-year vision and policy framework for Denver's growth"),
    (DC, DCM, "Planning & Zoning", "E-Permits (Online Building Permits)", "https://denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Community-Planning-and-Development/E-permits", "Online building and development permit submission"),
    (DC, DCM, "Planning & Zoning", "Building Permit Status", "https://www.denvergov.org/buildingpermitstatus", "Check status of submitted building and development permits"),
    # Floodplain / FEMA
    (DC, DCM, "Floodplain / FEMA", "Colorado Hazard Mapping Portal", "https://coloradohazardmapping.com/hazardMapping/floodplainMapping", "Statewide floodplain and hazard mapping resource"),
    (DC, DCM, "Floodplain / FEMA", "FEMA Map Service Center", "https://msc.fema.gov/", "Official FEMA FIRM maps by address"),
    # Economic Development
    (DC, DCM, "Economic Development", "Denver Economic Development & Opportunity", "https://denvergov.org/Government/Agencies-Departments-Offices/Agencies-Departments-Offices-Directory/Economic-Development-Opportunity", "Official economic development resources and business support"),
    (DC, DCM, "Economic Development", "Metro Denver EDC", "https://www.metrodenver.org/", "Regional economic development — nine-county metro area"),
    # Schools
    (DC, DCM, "Schools", "Denver Public Schools", "https://www.dpsk12.org/", "Primary district serving Denver County — boundary maps and school finder"),
]

for row_data in denver_rows:
    append_row(*row_data)
print(f"Denver County added. Total rows: {ws.max_row}")

# ── ARAPAHOE COUNTY ───────────────────────────────────────────────────────────
AC  = "Arapahoe County"
ACU = "Arapahoe County (Unincorporated)"

arapahoe_rows = [
    # Property Assessment
    (AC, ACU, "Property Assessment", "Arapahoe County Assessor", "https://www.arapahoeco.gov/your_county/county_departments/assessor/index.php", "Official assessor office — contact info, valuation, and resources"),
    (AC, ACU, "Property Assessment", "Property Search (Residential / Commercial)", "https://www.arapahoeco.gov/your_county/county_departments/assessor/property_search/search_residential_commercial_ag_and_vacant.php", "Online search for real estate assessment information"),
    (AC, ACU, "Property Assessment", "Parcel Search", "https://parcelsearch.arapahoegov.com/", "Interactive parcel search tool"),
    (AC, ACU, "Property Assessment", "Public Access NOW (Property Records)", "https://co-arapahoe.publicaccessnow.com/", "Assessor property records public database"),
    (AC, ACU, "Property Assessment", "NOV Appeals — Real Property", "https://www.arapahoeco.gov/your_county/county_departments/assessor/appeals/appealing_real_property_valuation.php", "Notice of Valuation appeal process and deadlines"),
    (AC, ACU, "Property Assessment", "Board of Equalization", "https://www.arapahoeco.gov/your_county/county_departments/assessor/board_of_equalization/index.php", "Second-level property valuation appeals"),
    (AC, ACU, "Property Assessment", "DOLA — Arapahoe County Assessor", "https://dpt.colorado.gov/locality/arapahoe-county-assessor", "State DPT profile page for Arapahoe County Assessor"),
    # Property Tax / Treasurer
    (AC, ACU, "Property Tax / Treasurer", "Treasurer Main Page", "https://www.arapahoeco.gov/your_county/county_departments/treasurer/index.php", "County treasurer — property tax info and services"),
    (AC, ACU, "Property Tax / Treasurer", "Property Tax Information", "https://www.arapahoeco.gov/your_county/county_departments/treasurer/property_tax_information/index.php", "Tax calculation, rates, and billing info"),
    (AC, ACU, "Property Tax / Treasurer", "Tax Search", "https://www.arapahoeco.gov/your_county/county_departments/treasurer/tax_search.php", "Search property tax accounts and values"),
    (AC, ACU, "Property Tax / Treasurer", "Pay Taxes Online", "https://www.arapahoeco.gov/your_county/county_departments/treasurer/pay_taxes/index.php", "Online property tax payment options"),
    # Clerk & Recorder
    (AC, ACU, "Clerk & Recorder / Official Records", "Official Record Search", "https://arapahoe.co.publicsearch.us/", "Deeds and recorded documents — public online search"),
    (AC, ACU, "Clerk & Recorder / Official Records", "Clerk & Recorder Office", "https://www.arapahoeco.gov/your_county/county_departments/clerk_and_recorder/index.php", "County clerk and recorder main page"),
    # GIS / Mapping
    (AC, ACU, "GIS / Mapping", "ArapaMAP (Interactive GIS)", "https://gis.arapahoegov.com/arapamap/", "Interactive GIS — parcel information and mapping"),
    (AC, ACU, "GIS / Mapping", "GIS Data Download", "https://gis.arapahoegov.com/datadownload/", "Download GIS data products and parcel information"),
    (AC, ACU, "GIS / Mapping", "Arapahoe Open Data Hub (ArcGIS)", "https://arapahoe-open-data-arapahoegov.hub.arcgis.com/", "County GIS data and mapping resources"),
    # Planning & Zoning
    (AC, ACU, "Planning & Zoning", "Planning and Land Development", "https://www.arapahoeco.gov/your_county/county_departments/public_works_and_development/divisions/planning_and_land_development/index.php", "Planning and development services for unincorporated county"),
    (AC, ACU, "Planning & Zoning", "Land Development Code", "https://www.arapahoeco.gov/your_county/county_departments/public_works_and_development/divisions/planning_and_land_development/land_development_code_and_regulations/land_development_code.php", "Official zoning regulations and development standards"),
    (AC, ACU, "Planning & Zoning", "Land Development Approvals", "https://www.arapahoeco.gov/your_county/county_departments/public_works_and_development/divisions/planning_and_land_development/land_development_approvals/index.php", "Development approval process and applications"),
    # Floodplain / FEMA
    (AC, ACU, "Floodplain / FEMA", "Arapahoe County Floodplain Management", "https://www.arapahoeco.gov/your_county/county_departments/public_works_and_development/divisions/engineering_services/floodplain_management.php", "County floodplain regulations and management"),
    (AC, ACU, "Floodplain / FEMA", "FEMA Map Service Center", "https://msc.fema.gov/", "Official FEMA FIRM maps by address"),
    (AC, ACU, "Floodplain / FEMA", "Colorado Hazard Mapping Portal", "https://coloradohazardmapping.com/hazardMapping/floodplainMapping", "Statewide floodplain and hazard mapping"),
    # Economic Development
    (AC, ACU, "Economic Development", "UACED — Unincorporated Arapahoe County EDC", "https://uaced.com/", "Economic development for unincorporated Arapahoe County"),
    (AC, ACU, "Economic Development", "Metro Denver — Arapahoe County", "https://www.metrodenver.org/do-business/communities/arapahoe", "Regional business development and site selection"),

    # Aurora
    (AC, "Aurora", "Property Assessment", "Aurora Property Info System", "https://ags.auroragov.org/propinfo/propinfo.html", "Interactive parcel lookup for Aurora properties"),
    (AC, "Aurora", "Planning & Zoning", "Aurora Planning & Development Services", "https://www.auroragov.org/city_hall/departments/planning___development_services", "City planning and development services department"),
    (AC, "Aurora", "Planning & Zoning", "Aurora Building & Zoning Code", "https://www.auroragov.org/business_services/development_center/codes_rules/building___zoning_code", "City development regulations and codes"),
    (AC, "Aurora", "Planning & Zoning", "Aurora Planning & Zoning Maps", "https://www.auroragov.org/city_hall/maps/planning_and_zoning_maps", "City planning and zoning map viewer"),
    (AC, "Aurora", "Economic Development", "Aurora Economic Development Council", "https://auroraedc.com/", "Private EDC for the Aurora region"),

    # Littleton
    (AC, "Littleton", "Planning & Zoning", "Littleton ULUC & Zoning Portal", "https://www.littletonco.gov/Building-Development/Land-Planning-Entitlement/Plans-and-Regulations/ULUC-and-Zoning-Portal", "Interactive zoning map and unified land use code"),
    (AC, "Littleton", "Planning & Zoning", "Littleton Plans & Regulations", "https://www.littletonco.gov/Building-Development/Land-Planning-Entitlement/Plans-and-Regulations", "Comprehensive planning documents and development standards"),
    (AC, "Littleton", "GIS / Mapping", "Littleton Open GIS Data Hub", "https://data-littleton.opendata.arcgis.com/", "Open data hub with GIS data downloads"),
    (AC, "Littleton", "Economic Development", "Littleton Economic Development", "https://www.littletonco.gov/Business/Economic-Development", "City economic development resources"),
    (AC, "Littleton", "Schools", "Littleton Public Schools", "https://www.littletonpublicschools.net/", "Primary school district serving Littleton"),

    # Englewood
    (AC, "Englewood", "Planning & Zoning", "Englewood Zoning & Development", "https://www.englewoodco.gov/government/city-departments/community-development/zoning", "City zoning regulations and development standards"),
    (AC, "Englewood", "Planning & Zoning", "Englewood Zoning Map", "https://www.englewoodco.gov/government/city-departments/community-development/zoning/zoning-map", "Interactive zoning map for Englewood"),
    (AC, "Englewood", "Planning & Zoning", "Englewood Comprehensive Plan", "https://www.englewoodco.gov/government/city-departments/community-development/englewood-comprehensive-plan", "City comprehensive development plan"),
    (AC, "Englewood", "Schools", "Englewood Schools", "https://www.englewoodschools.net/", "Primary school district serving Englewood"),

    # Centennial
    (AC, "Centennial", "Property Assessment", "Centennial Property Search", "https://www.centennialco.gov/Online-Services/Property-Search", "City property search tool"),
    (AC, "Centennial", "Planning & Zoning", "Centennial Land Development Code", "https://www.centennialco.gov/Government/Departments/Community-Development/Land-Development-Code", "City zoning and development regulations"),
    (AC, "Centennial", "GIS / Mapping", "Centennial Open Data Hub", "https://open-centennial.opendata.arcgis.com/", "City open data portal with mapping and parcel data"),
    (AC, "Centennial", "Economic Development", "Centennial Economic Development", "https://www.centennialco.gov/Government/Departments/Economic-Development", "City economic development services and incentives"),

    # Sheridan
    (AC, "Sheridan", "Planning & Zoning", "Sheridan Planning & Zoning", "https://ci.sheridan.co.us/443/Planning-Zoning", "City planning and zoning information"),
    (AC, "Sheridan", "Planning & Zoning", "Sheridan Municipal Code", "https://library.municode.com/co/sheridan", "City complete municipal code and ordinances"),
    (AC, "Sheridan", "Schools", "Sheridan School District 2", "https://www.ssd2.org/", "Primary school district serving Sheridan"),

    # Cherry Hills Village
    (AC, "Cherry Hills Village", "Planning & Zoning", "Cherry Hills Village Planning & Zoning", "https://www.cherryhillsvillage.com/382/Planning-Zoning-Division", "City planning and zoning services"),
    (AC, "Cherry Hills Village", "Planning & Zoning", "Cherry Hills Village Zoning Maps", "https://www.cherryhillsvillage.com/392/Zoning-District-Maps", "City zoning district maps and information"),
    (AC, "Cherry Hills Village", "Schools", "Cherry Creek School District", "https://www.cherrycreekschools.org/", "Primary school district serving Cherry Hills Village"),

    # Greenwood Village
    (AC, "Greenwood Village", "Planning & Zoning", "Greenwood Village Planning & Development Review", "https://www.greenwoodvillage.com/1269/Planning-Development-Review", "City planning and development review process"),
    (AC, "Greenwood Village", "Planning & Zoning", "Greenwood Village Zoning Code", "https://library.municode.com/co/greenwood_village/codes/municipal_code?nodeId=CH16LADECO", "Complete city land development code"),
    (AC, "Greenwood Village", "Schools", "Cherry Creek School District", "https://www.cherrycreekschools.org/", "Primary school district serving Greenwood Village"),

    # Glendale
    (AC, "Glendale", "Planning & Zoning", "Glendale Planning & Development", "https://www.glendale.co.us/150/Planning-Development", "City planning and development services"),
    (AC, "Glendale", "Schools", "Cherry Creek School District", "https://www.cherrycreekschools.org/", "Primary school district serving Glendale"),
]

for row_data in arapahoe_rows:
    append_row(*row_data)
print(f"Arapahoe County added. Total rows: {ws.max_row}")

# ── DOUGLAS COUNTY ────────────────────────────────────────────────────────────
DGC  = "Douglas County"
DGCU = "Douglas County (Unincorporated)"

douglas_rows = [
    # Property Assessment
    (DGC, DGCU, "Property Assessment", "Assessor Property Search", "https://apps.douglas.co.us/assessor/web/", "Property search by address, owner, account, or parcel number"),
    (DGC, DGCU, "Property Assessment", "Assessor Advanced Search", "https://apps.douglas.co.us/assessor/advanced-search/", "Generate reports — ownership, valuation, and improvements data"),
    (DGC, DGCU, "Property Assessment", "Douglas County Assessor Main", "https://www.douglasco.gov/assessor/", "Assessment process, valuations, appeals, and tax information"),
    (DGC, DGCU, "Property Assessment", "Commercial Property Assessment", "https://www.douglas.co.us/assessor/commercial-property/", "Commercial and industrial property assessment information"),
    (DGC, DGCU, "Property Assessment", "Online Appeal System", "https://appeals.spatialest.com/co-douglas", "Online portal for filing property valuation appeals"),
    (DGC, DGCU, "Property Assessment", "DOLA — Douglas County Assessor", "https://dpt.colorado.gov/locality/douglas-county-assessor", "State DPT profile page for Douglas County Assessor"),
    # Property Tax / Treasurer
    (DGC, DGCU, "Property Tax / Treasurer", "Tax Account Search / EagleWeb", "https://apps.douglas.co.us/treasurer/web/", "View tax accounts and make online payments"),
    (DGC, DGCU, "Property Tax / Treasurer", "Treasurer Department", "https://www.douglas.co.us/treasurer/", "General treasurer office info and services"),
    (DGC, DGCU, "Property Tax / Treasurer", "Property Tax Payment Info", "https://www.douglasco.gov/treasurer/about-property-tax/", "Tax payment options, due dates, mill levy info"),
    (DGC, DGCU, "Property Tax / Treasurer", "Your Douglas County Taxes", "https://www.yourdougcotaxes.com/", "Tax breakdown by taxing authority"),
    # Clerk & Recorder
    (DGC, DGCU, "Clerk & Recorder / Official Records", "Landmark Web Portal (Recorded Documents)", "https://apps.douglas.co.us/LandmarkWeb", "Deeds, mortgages, and recorded documents online search"),
    (DGC, DGCU, "Clerk & Recorder / Official Records", "Clerk & Recorder Main", "https://www.douglasco.gov/clerk-recorder/", "County clerk and recorder office info"),
    (DGC, DGCU, "Clerk & Recorder / Official Records", "Recording Documents", "https://www.douglas.co.us/recording/recording-documents/", "Instructions and forms for recording documents"),
    # GIS / Mapping
    (DGC, DGCU, "GIS / Mapping", "DougCo Hub (ArcGIS Open Data)", "https://dcdata-dougco.opendata.arcgis.com/", "County GIS data, maps, and downloadable datasets"),
    (DGC, DGCU, "GIS / Mapping", "DC Maps — Interactive Zoning Map", "https://apps.douglas.co.us/dcmaps/map.html?mapInstance=zoning", "Interactive map to determine property zoning"),
    (DGC, DGCU, "GIS / Mapping", "GIS Maps and Apps", "https://www.douglas.co.us/information-technology/gis-maps-apps/", "County GIS resources and applications directory"),
    # Planning & Zoning
    (DGC, DGCU, "Planning & Zoning", "Planning Services", "https://www.douglas.co.us/planning/", "General planning department info and services"),
    (DGC, DGCU, "Planning & Zoning", "Zoning Resolution", "https://www.douglas.co.us/planning/development-review-regulations/zoning/", "Full text of Douglas County Zoning Resolution"),
    (DGC, DGCU, "Planning & Zoning", "Douglas County 2040 Comprehensive Master Plan", "https://www.douglas.co.us/documents/cmp-toc.pdf/", "County-wide comprehensive master plan (2040 horizon)"),
    (DGC, DGCU, "Planning & Zoning", "Building Permits Portal", "https://apps.douglas.co.us/building/services/", "Permit search, applications, and inspections"),
    # Floodplain / FEMA
    (DGC, DGCU, "Floodplain / FEMA", "Douglas County Floodplain Info", "https://www.douglas.co.us/public-works/stormwater/flood-plain-information/", "Floodplain info, FIRM maps, and permit requirements"),
    (DGC, DGCU, "Floodplain / FEMA", "FEMA Map Service Center", "https://msc.fema.gov/", "Official FEMA FIRM maps by address"),
    (DGC, DGCU, "Floodplain / FEMA", "Colorado Hazard Mapping Portal", "https://coloradohazardmapping.com/hazardMapping/floodplainMapping", "Statewide floodplain and hazard mapping"),
    # Economic Development
    (DGC, DGCU, "Economic Development", "Douglas County EDC", "https://douglascountyedc.com/", "Primary economic development agency for Douglas County"),
    (DGC, DGCU, "Economic Development", "Economic Development — Planning", "https://www.douglas.co.us/planning/economic-development/", "County planning for economic development initiatives"),
    # Schools
    (DGC, DGCU, "Schools", "Douglas County School District RE-1", "https://www.dcsdk12.org/", "Primary district serving most of Douglas County"),

    # Highlands Ranch (unincorporated)
    (DGC, "Highlands Ranch", "Planning & Zoning", "Highlands Ranch Metro District Development", "https://www.highlandsranch.org/government/public-works/development", "Metro district development approval procedures"),
    (DGC, "Highlands Ranch", "Planning & Zoning", "Highlands Ranch Planned Development Guide", "https://www.douglas.co.us/documents/highlands-ranch-pd-summary.pdf/", "Planned Community District development guidelines"),

    # Castle Rock
    (DGC, "Castle Rock", "Property Assessment", "Castle Rock Property Tax Address Search", "https://www.crgov.com/2415/Property-Tax-Address-Search", "Town-hosted tax assessment lookup tool"),
    (DGC, "Castle Rock", "Planning & Zoning", "Castle Rock Planning Division", "https://www.crgov.com/1888/Planning", "Town planning services and development review"),
    (DGC, "Castle Rock", "Planning & Zoning", "Castle Rock Zoning Division", "https://www.crgov.com/1889/Zoning", "Zoning administration and information"),
    (DGC, "Castle Rock", "Planning & Zoning", "Castle Rock Municipal Code — Title 17 Zoning", "https://library.municode.com/co/castle_rock/codes/municipal_code?nodeId=TIT17ZO", "Zoning ordinance and regulations"),
    (DGC, "Castle Rock", "Planning & Zoning", "Castle Rock Vision & Master Plan", "https://www.crgov.com/2442/Vision-and-Master-Plan", "Town 2030 Comprehensive Master Plan"),
    (DGC, "Castle Rock", "Planning & Zoning", "Castle Rock Building Division", "https://www.crgov.com/1887/Building", "Building permits, inspections, and code enforcement"),
    (DGC, "Castle Rock", "GIS / Mapping", "Castle Rock Maps & GIS", "https://www.crgov.com/2295/Maps-GIS", "Town maps and GIS resources"),

    # Parker
    (DGC, "Parker", "Planning & Zoning", "Parker Land Development Ordinance", "https://library.municode.com/co/parker/codes/municipal_code?nodeId=TIT13LADEOR", "Complete land development and zoning code"),
    (DGC, "Parker", "Planning & Zoning", "Parker 2035 Master Plan", "https://parkerco.gov/542/Parker-2035-Master-Plan", "Town comprehensive master plan"),
    (DGC, "Parker", "Planning & Zoning", "Parker Building Division", "https://parkerco.gov/114/Building-Division", "Building permits, inspections, and code enforcement"),
    (DGC, "Parker", "Planning & Zoning", "Parker eTRAKiT Permit Portal", "https://prkc-trk.aspgov.com/eTRAKiT/", "Online permit application and tracking system"),

    # Lone Tree
    (DGC, "Lone Tree", "Planning & Zoning", "Lone Tree Planning Division", "https://cityoflonetree.com/departments/community-development/planning/", "City planning services and development review"),
    (DGC, "Lone Tree", "Planning & Zoning", "Lone Tree Elevated Comprehensive Plan", "https://cityoflonetree.com/community-development/regulations-and-master-plans/lone-tree-elevated/", "Updated comprehensive plan for the city"),
    (DGC, "Lone Tree", "Planning & Zoning", "Lone Tree Building Permits (Accela)", "https://aca-prod.accela.com/lonetree/default.aspx", "Online building permit portal"),

    # Castle Pines
    (DGC, "Castle Pines", "Planning & Zoning", "Castle Pines Land Use & Zoning", "https://www.castlepinesco.gov/city-services/city-departments/community-development/land-use-zoning/", "Zoning and land use regulations"),
    (DGC, "Castle Pines", "Planning & Zoning", "Castle Pines Comprehensive Plan", "https://www.castlepinesco.gov/comprehensive-plan/", "City comprehensive master plan"),
    (DGC, "Castle Pines", "Planning & Zoning", "Castle Pines Building Department", "https://www.castlepinesco.gov/city-services/city-departments/community-development/building-division/", "Building permits and inspections"),

    # Larkspur
    (DGC, "Larkspur", "Planning & Zoning", "Larkspur Planning & Development", "https://www.townoflarkspur.org/planning-development", "Town planning services and information"),
    (DGC, "Larkspur", "Planning & Zoning", "Larkspur Municipal Code", "https://library.municode.com/co/larkspur", "Town municipal code including zoning"),
]

for row_data in douglas_rows:
    append_row(*row_data)
print(f"Douglas County added. Total rows: {ws.max_row}")

# ── UPDATE COUNTY DROPDOWN ───────────────────────────────────────────────────
ws.data_validations.dataValidation = [
    dv for dv in ws.data_validations.dataValidation
    if dv.sqref != "D3"
]

dv = DataValidation(
    type="list",
    formula1='"Adams County,Arapahoe County,Boulder County,Broomfield County,Denver County,Douglas County,Jefferson County,Larimer County,Weld County"',
    showDropDown=False,
    showErrorMessage=True,
    errorTitle="Invalid County",
    error="Please select a county from the list."
)
dv.sqref = "D3"
ws.add_data_validation(dv)
print("County dropdown updated to 9 counties.")

wb.save(DST)
print("Saved.")
