"""
Add Jefferson County and Larimer County to the Resources sheet.
Update county dropdown to include all 6 counties.
"""

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

DST = "/home/user/Appraisal-Template/Analysis_Workbook_OPTIMIZED.xlsm"
wb = load_workbook(DST, keep_vba=True)
ws = wb["Resources"]

URL_FONT = Font(color="0563C1", underline="single", size=10)

def prio(r):
    return (
        f'=IF(B{r}<>$D$3,"◌ Other",'
        f'IF(C{r}="REGIONAL / ALL MUNICIPALITIES","Regional",'
        f'IF(OR('
        f'ISNUMBER(SEARCH(LEFT(\'Property Info\'!$B$17,MAX(1,FIND(",",\'Property Info\'!$B$17&",")-1)),C{r})),'
        f'ISNUMBER(SEARCH(\'Property Info\'!$B$18,C{r}))'
        f')),"★ Active","–")))'
    )

def append_row(county, municipality, category, name, url, description):
    r = ws.max_row + 1
    ws.cell(r, 1, prio(r))
    ws.cell(r, 2, county)
    ws.cell(r, 3, municipality)
    ws.cell(r, 4, category)
    ws.cell(r, 5, name)
    uc = ws.cell(r, 6, url)
    uc.hyperlink = url
    uc.font = URL_FONT
    ws.cell(r, 7, description)
    ws.row_dimensions[r].height = 15

# ── JEFFERSON COUNTY ─────────────────────────────────────────────────────────
JC  = "Jefferson County"
JCU = "Jefferson County (Unincorporated)"

jefferson_rows = [
    # Property Assessment
    (JC, JCU, "Property Assessment", "Assessor Property Records Search", "https://propertysearch.jeffco.us/", "Search by address, parcel ID, owner name, or schedule number; includes valuation history and sales data"),
    (JC, JCU, "Property Assessment", "Assessor Office", "https://www.jeffco.us/658/Assessor", "County assessor homepage — contact info, appointment booking, and property search links"),
    (JC, JCU, "Property Assessment", "Property Valuation & Appeals", "https://www.jeffco.us/4593/20252026-Property-Valuation-and-Appeal-I", "NOV information, appeal deadlines, and assessment appeal process"),
    (JC, JCU, "Property Assessment", "Appealing Real Property Valuation", "https://www.jeffco.us/435/Appealing-Real-Property-Valuation", "Instructions and timeline for appealing assessed values to the BOE"),
    (JC, JCU, "Property Assessment", "Commercial Property Valuation", "https://www.jeffco.us/514/Commercial-Property", "Commercial valuation methods, Marshall Valuation data, and appraiser contacts"),
    (JC, JCU, "Property Assessment", "Parcel Maps", "https://www.jeffco.us/713/Parcel-Maps", "Interactive parcel mapping with graphical selection, printing, and query tools"),
    # Property Tax / Treasurer
    (JC, JCU, "Property Tax / Treasurer", "Treasurer Property Records Search", "https://treasurerpropertysearch.jeffco.us/", "Tax information, bills, and account search by address, PIN, or owner"),
    (JC, JCU, "Property Tax / Treasurer", "Treasurer Office", "https://www.jeffco.us/809/Treasurer", "Treasurer contact info, payment options, and office hours"),
    (JC, JCU, "Property Tax / Treasurer", "Property Taxes", "https://www.jeffco.us/2415/Property-Taxes", "Mill levy, tax determination process, and payment details"),
    (JC, JCU, "Property Tax / Treasurer", "Mill Levy & Tax Rate Determination", "https://www.jeffco.us/822/Determination-of-Tax-Rates", "How tax rates are determined and mill levy information by district"),
    # Clerk & Recorder
    (JC, JCU, "Clerk & Recorder / Official Records", "Clerk & Recorder Web Access (Land Records)", "https://landrecords.co.jefferson.co.us/", "Deeds, liens, plats, and UCC filings back to 1859"),
    (JC, JCU, "Clerk & Recorder / Official Records", "Clerk & Recorder Office", "https://www.jeffco.us/4871/Clerk-Recorder", "County clerk and recorder office information and services"),
    (JC, JCU, "Clerk & Recorder / Official Records", "Recording Services", "https://www.jeffco.us/398/Recording", "Document recording procedures and submission information"),
    # GIS / Mapping
    (JC, JCU, "GIS / Mapping", "Jefferson County Web Maps (jMap)", "https://gis.jeffco.us/", "Interactive GIS portal with property, zoning, development, and tax layers"),
    (JC, JCU, "GIS / Mapping", "Jefferson County Open Data (ArcGIS Hub)", "https://data-jeffersoncounty.opendata.arcgis.com/", "Downloadable GIS datasets: parcels, zoning, infrastructure, and more"),
    (JC, JCU, "GIS / Mapping", "Maps & Data Download", "https://www.jeffco.us/3165/Maps-Data-Download", "County GIS data downloads including parcel and property boundary files"),
    # Planning & Zoning
    (JC, JCU, "Planning & Zoning", "Planning & Zoning Division", "https://www.jeffco.us/786/Planning-Zoning", "County planning and zoning office — development review and land use regulations"),
    (JC, JCU, "Planning & Zoning", "Unified Land Use Code (ULUC)", "https://togetherjeffco.com/uluc", "Consolidated land use code with zoning, development standards, and design criteria"),
    (JC, JCU, "Planning & Zoning", "Zoning Resolution", "https://www.jeffco.us/2460/Zoning-Resolution", "District designations, uses, setbacks, and lot size requirements"),
    (JC, JCU, "Planning & Zoning", "Comprehensive Master Plan", "https://www.jeffco.us/2468/Comprehensive-Master-Plan", "County comprehensive plan guiding long-range land use decisions"),
    (JC, JCU, "Planning & Zoning", "Area Plans", "https://www.jeffco.us/2477/Area-Plans", "Central Mountains, Conifer/285 Corridor, and plains area-specific plans"),
    (JC, JCU, "Planning & Zoning", "Online Permitting & Inspection Portal", "https://www.jeffco.us/2200/Online-Permitting-Inspection-Requests", "Online permit submission and inspection request via Jeffco Citizen Portal"),
    # Floodplain / FEMA
    (JC, JCU, "Floodplain / FEMA", "Floodplain Management", "https://www.jeffco.us/2695/Floodplain-Management", "Floodplain regulations, development permits, and management info"),
    (JC, JCU, "Floodplain / FEMA", "FEMA Map Service Center", "https://msc.fema.gov/", "Official FEMA FIRM maps by address"),
    (JC, JCU, "Floodplain / FEMA", "Colorado Hazard Mapping Portal", "https://coloradohazardmapping.com/hazardMapping/floodplainMapping", "Statewide floodplain and hazard mapping resource"),
    # Economic Development
    (JC, JCU, "Economic Development", "Jefferson County EDC", "https://jeffcoedc.org/", "Business attraction, expansion, and retention in Jeffco"),
    (JC, JCU, "Economic Development", "Business Incentives & Resources", "https://jeffcoedc.org/resources-and-incentives/", "Enterprise Zone credits, job tax credits, and business incentive programs"),
    # Schools
    (JC, JCU, "Schools", "Jefferson County School District R-1", "https://www.jeffco.k12.co.us/", "Primary county district — ~81,000 students in 160 schools"),
    (JC, JCU, "Schools", "Littleton Public Schools (District 6)", "https://www.littletonpublicschools.net/", "Serves portions of southern Jeffco near Littleton"),
    # DOLA
    (JC, JCU, "Property Assessment", "DOLA — Jefferson County Assessor", "https://dpt.colorado.gov/locality/jefferson-county-assessor", "State DPT profile page for Jefferson County Assessor"),

    # Golden
    (JC, "Golden", "Planning & Zoning", "Golden Land Use & Development", "https://www.cityofgolden.gov/business/land_use_development/index.php", "Planning, zoning, and development review for Golden"),
    (JC, "Golden", "Planning & Zoning", "Golden Municipal Code Title 18 — Zoning", "https://library.municode.com/co/golden/codes/municipal_code?nodeId=TIT18PLZO", "Golden zoning districts, uses, and development standards"),
    (JC, "Golden", "Planning & Zoning", "Golden Development Applications", "https://www.cityofgolden.gov/business/land_use_development/development_applications.php", "Development application process and required documents"),
    (JC, "Golden", "Planning & Zoning", "Guiding Golden Comprehensive Plan 2026", "https://www.guidinggolden.com/comprehensive-plan", "City comprehensive plan adopted 2026 guiding future development"),

    # Lakewood
    (JC, "Lakewood", "Planning & Zoning", "Lakewood Zoning & Standards", "https://www.lakewoodco.gov/Local-Government/Departments/Sustainability-and-Community-Development/Zoning-Standards", "City of Lakewood zoning information and development standards"),
    (JC, "Lakewood", "Planning & Zoning", "Envision Lakewood 2040 Comprehensive Plan", "https://www.lakewoodco.gov/Local-Government/Departments/Sustainability-and-Community-Development/Comprehensive-Planning-Main/Envision-Lakewood-2040-Comprehensive-Plan", "Comprehensive plan adopted 2025 guiding development through 2040"),
    (JC, "Lakewood", "GIS / Mapping", "Lakewood Property Search & GIS", "https://propertysearch.jeffco.us/", "Jefferson County assessor search for Lakewood properties"),

    # Arvada
    (JC, "Arvada", "Planning & Zoning", "Arvada Zoning", "https://www.arvadaco.gov/168/Zoning", "Zoning district information and verification for Arvada properties"),
    (JC, "Arvada", "Planning & Zoning", "Arvada Land Development Code", "https://www.arvadaco.gov/179/Land-Development-Code", "Zoning districts, land uses, and development standards"),
    (JC, "Arvada", "Planning & Zoning", "Arvada Comprehensive Plan Update 2026", "https://www.arvadaco.gov/1374/2026-27-Comprehensive-Plan", "Arvada 10-year comprehensive planning process"),
    (JC, "Arvada", "GIS / Mapping", "City of Arvada Open Data", "https://gis-arvada.opendata.arcgis.com/", "Arvada ArcGIS open data hub with downloadable spatial datasets"),

    # Wheat Ridge
    (JC, "Wheat Ridge", "Planning & Zoning", "Wheat Ridge Zoning & Development Code", "https://www.ci.wheatridge.co.us/1114/Zoning-Development-Code", "Chapter 26 — zoning code and development regulations"),
    (JC, "Wheat Ridge", "Planning & Zoning", "Find Your Zoning — Wheat Ridge", "https://www.ci.wheatridge.co.us/384/Find-Your-Zoning", "Tool to identify property zoning for Wheat Ridge locations"),
    (JC, "Wheat Ridge", "Planning & Zoning", "Wheat Ridge Land Use Planning", "https://www.ci.wheatridge.co.us/394/Land-Use-Planning", "Development review and land use applications"),

    # Edgewater
    (JC, "Edgewater", "Planning & Zoning", "Edgewater Zoning Code (Ch. 16)", "https://library.municode.com/co/edgewater/codes/municipal_code?nodeId=EDMUCO_CH16ZO", "Edgewater Municipal Code Chapter 16 — zoning districts and regulations"),
    (JC, "Edgewater", "Planning & Zoning", "Edgewater Planning & Zoning Commission", "https://www.edgewaterco.gov/your-government/boards-commissions/planning-zoning-commission", "Planning and zoning commission information"),

    # Morrison
    (JC, "Morrison", "Planning & Zoning", "Town of Morrison", "https://www.town.morrison.co.us/", "Town of Morrison official website — planning and development resources"),

    # Littleton (Jeffco portion)
    (JC, "Littleton", "Planning & Zoning", "Littleton Land Use Code & Zoning Portal", "https://www.littletonco.gov/Building-Development/Land-Planning-Entitlement/Plans-and-Regulations/ULUC-and-Zoning-Portal", "Littleton city zoning and unified land use code"),
]

for row_data in jefferson_rows:
    append_row(*row_data)

print(f"Jefferson County rows added. Total rows: {ws.max_row}")

# ── LARIMER COUNTY ────────────────────────────────────────────────────────────
LC  = "Larimer County"
LCU = "Larimer County (Unincorporated)"

larimer_rows = [
    # Property Assessment
    (LC, LCU, "Property Assessment", "Assessor Property Search", "https://www.larimer.gov/assessor/search", "Search by address, owner name, or parcel number; access ownership, valuations, land/building details, and tax data"),
    (LC, LCU, "Property Assessment", "Assessor Office", "https://www.larimer.gov/assessor", "Larimer County Assessor main page — 200 W. Oak Street, Fort Collins"),
    (LC, LCU, "Property Assessment", "Assessor Public Data Center", "https://www.larimer.gov/assessor/publicdata", "GIS digital data portal with publicly available assessor data downloads"),
    (LC, LCU, "Property Assessment", "Property Valuation Appeals", "https://www.larimer.gov/bocc/property-valuation-appeals", "NOV appeal process and deadlines for residential and non-residential properties"),
    # Property Tax / Treasurer
    (LC, LCU, "Property Tax / Treasurer", "Property Tax Search", "https://www.larimer.gov/treasurer/search", "Search tax info by account, parcel, owner name, address, or tax year"),
    (LC, LCU, "Property Tax / Treasurer", "Pay Property Taxes Online", "https://www.larimer.gov/treasurer/pay", "Secure online payment by eCheck or credit/debit card"),
    (LC, LCU, "Property Tax / Treasurer", "Treasurer & Public Trustee Office", "https://www.larimer.gov/treasurer", "Treasurer main office — 200 W. Oak Street, Suite 2100, Fort Collins"),
    # Clerk & Recorder
    (LC, LCU, "Clerk & Recorder / Official Records", "Easy Access Online Records Portal", "https://records.larimer.org/landmarkweb/", "Recorded documents from 1939 to present — search by name, legal description, or document type"),
    (LC, LCU, "Clerk & Recorder / Official Records", "Clerk & Recorder Office", "https://www.larimer.gov/clerk", "Main clerk and recorder office — 200 W. Oak Street, First Floor, Fort Collins"),
    (LC, LCU, "Clerk & Recorder / Official Records", "Recording Documents", "https://www.larimer.gov/clerk/recording/docs", "Document recording procedures and certified copies"),
    # GIS / Mapping
    (LC, LCU, "GIS / Mapping", "Enterprise GIS Portal (Geocortex)", "https://maps1.larimer.org/gvh/?Viewer=LIL", "Interactive GIS with parcel info, zoning, flood data, and aerial imagery"),
    (LC, LCU, "GIS / Mapping", "GIS Digital Data Download", "https://www.larimer.gov/it/gis/digital-data", "Downloadable GIS datasets — parcels, contours, LiDAR, and more"),
    (LC, LCU, "GIS / Mapping", "GIS Information Hub", "https://www.larimer.gov/it/gis", "Main GIS portal with mapping resources, data access, and web services"),
    # Planning & Zoning
    (LC, LCU, "Planning & Zoning", "Planning Division", "https://www.larimer.gov/planning", "County planning — 200 W. Oak Street, Third Floor, Fort Collins"),
    (LC, LCU, "Planning & Zoning", "Land Use Code", "https://www.larimer.gov/planning/land-use-code", "Larimer County Land Use Code (effective December 8, 2025) — zoning districts, overlays, and development standards"),
    (LC, LCU, "Planning & Zoning", "Zoning Districts", "https://www.larimer.gov/planning/zoning/districts", "Unincorporated Larimer County zoning districts, minimum lot sizes, and setbacks"),
    (LC, LCU, "Planning & Zoning", "Zoning Parcel Search", "https://www.larimer.gov/planning/zoning/search", "Look up zoning district for any parcel in unincorporated Larimer County"),
    # Floodplain / FEMA
    (LC, LCU, "Floodplain / FEMA", "Floodplains & Stormwater", "https://www.larimer.gov/engineering/floodplains-and-stormwater/floodplains", "Interactive floodplain map — FEMA, Best Available, and Flood-Prone Areas"),
    (LC, LCU, "Floodplain / FEMA", "FEMA Map Service Center", "https://msc.fema.gov/", "Official FEMA FIRM maps by address"),
    (LC, LCU, "Floodplain / FEMA", "Colorado Hazard Mapping Portal", "https://coloradohazardmapping.com/hazardMapping/floodplainMapping", "Statewide floodplain and hazard mapping resource"),
    # Economic Development
    (LC, LCU, "Economic Development", "Larimer County Economic & Workforce Development", "https://www.larimer.gov/ewd", "Business assistance, workforce solutions, and incentive programs"),
    (LC, LCU, "Economic Development", "Small Business Development Center — Larimer", "https://sbdc.colorado.gov/larimer", "Free confidential consulting and workshops for Larimer County businesses"),
    # Schools
    (LC, LCU, "Schools", "Poudre School District R-1", "https://www.psdschools.org/", "Serves Fort Collins, Wellington, Timnath, and unincorporated Larimer County — 52 schools"),
    (LC, LCU, "Schools", "Thompson School District R2-J", "https://www.tsd.org/", "Serves Loveland, Berthoud, and portions of Larimer/Weld/Boulder counties — 32 schools"),
    (LC, LCU, "Schools", "Estes Park School District R-3", "https://www.estesschools.org/", "Serves Estes Park area — 3 schools"),
    # DOLA
    (LC, LCU, "Property Assessment", "DOLA — Larimer County Assessor", "https://dpt.colorado.gov/locality/larimer-county-assessor", "State DPT profile page for Larimer County Assessor"),

    # Fort Collins
    (LC, "Fort Collins", "Planning & Zoning", "Fort Collins Land Use Code", "https://www.fcgov.com/planning-development-services/luc", "City Land Use Code with zoning regulations (effective May 17, 2024)"),
    (LC, "Fort Collins", "Planning & Zoning", "Fort Collins Planning Services", "https://www.fcgov.com/planning/", "Zoning info, City Plan (comprehensive plan), and development services"),
    (LC, "Fort Collins", "Planning & Zoning", "Fort Collins Zoning Information", "https://www.fcgov.com/zoning", "Zoning districts, zones, and permitted uses in Fort Collins"),
    (LC, "Fort Collins", "GIS / Mapping", "Fort Collins GIS Portal", "https://gisweb.fcgov.com/HTML5Viewer/Index.html?Viewer=FCMaps", "City of Fort Collins GIS mapping and data access portal"),
    (LC, "Fort Collins", "Economic Development", "Fort Collins Economic Health Office", "https://www.fcgov.com/business/", "Business resources and development services"),
    (LC, "Fort Collins", "Economic Development", "Fort Collins Area Chamber — Economic Development", "https://fortcollinschamber.com/resources/economic-development/", "Chamber economic development resources and business support"),

    # Loveland
    (LC, "Loveland", "Planning & Zoning", "Loveland Unified Development Code (Title 18)", "https://online.encodeplus.com/regs/loveland-co/", "Zoning, subdivision, and annexation regulations"),
    (LC, "Loveland", "Planning & Zoning", "Loveland Current Planning Division", "https://www.lovgov.org/services/development-services/current-planning", "Planning services and development review"),
    (LC, "Loveland", "Planning & Zoning", "Loveland Property Zoning Search", "https://www.lovgov.org/business/view-property-information", "Search zoning information for Loveland properties"),
    (LC, "Loveland", "Floodplain / FEMA", "Loveland Flood Management", "https://www.lovgov.org/services/public-works/stormwater/flood-management", "City floodplain info and management programs"),
    (LC, "Loveland", "Economic Development", "Loveland Economic Development", "https://lovelandeconomicdevelopment.org/", "City economic development services and business growth support"),

    # Estes Park
    (LC, "Estes Park", "Planning & Zoning", "Estes Park Development Code & Comprehensive Plan", "https://estespark.colorado.gov/developmentcode", "Development Code (updated 2025-26) and Estes Forward Comprehensive Plan (2022)"),
    (LC, "Estes Park", "Planning & Zoning", "Estes Park Planning & Zoning", "https://estespark.colorado.gov/planningzoning", "Planning and zoning services, development consultations"),
    (LC, "Estes Park", "Planning & Zoning", "Estes Park Maps & Zoning Portal", "https://estespark.colorado.gov/maps", "Interactive mapping for zoning districts and planning information"),
    (LC, "Estes Park", "Floodplain / FEMA", "Estes Park Floodplain Management", "https://estespark.colorado.gov/floodplains", "Floodplain management, hydrology studies, and flood risk information"),

    # Berthoud
    (LC, "Berthoud", "Planning & Zoning", "Berthoud Planning Department", "https://www.berthoud.org/178/Planning-Department", "Town planning and development services"),
    (LC, "Berthoud", "Planning & Zoning", "Berthoud Development Code (Title 30)", "https://library.municode.com/co/berthoud/codes/code_of_ordinances?nodeId=CH30BEDECO", "Zoning districts, regulations, and land use standards"),

    # Windsor (Larimer portion)
    (LC, "Windsor", "Planning & Zoning", "Windsor Planning & Zoning", "https://www.windsorgov.com/1229/Planning-Zoning", "Town planning, zoning, and development review (also serves Weld County portion)"),

    # Timnath
    (LC, "Timnath", "Planning & Zoning", "Timnath Planning Division", "https://timnath.org/planning/", "Town of Timnath planning services and development review"),
    (LC, "Timnath", "Planning & Zoning", "Timnath Land Use Code", "https://timnath.org/town-ordinances/", "Land Use Code (Version 22, effective March 2, 2026)"),

    # Wellington
    (LC, "Wellington", "Planning & Zoning", "Wellington Building & Planning", "https://www.wellingtoncolorado.gov/150/Departments", "Town building and planning services"),
    (LC, "Wellington", "Planning & Zoning", "Wellington Planning Commission", "https://www.wellingtoncolorado.gov/160/Planning-Commission", "Planning Commission responsible for comprehensive planning and land use"),
]

for row_data in larimer_rows:
    append_row(*row_data)

print(f"Larimer County rows added. Total rows: {ws.max_row}")

# ── UPDATE COUNTY DROPDOWN ───────────────────────────────────────────────────
ws.data_validations.dataValidation = [
    dv for dv in ws.data_validations.dataValidation
    if dv.sqref != "D3"
]

dv = DataValidation(
    type="list",
    formula1='"Adams County,Boulder County,Broomfield County,Jefferson County,Larimer County,Weld County"',
    showDropDown=False,
    showErrorMessage=True,
    errorTitle="Invalid County",
    error="Please select a county from the list."
)
dv.sqref = "D3"
ws.add_data_validation(dv)
print("County dropdown updated to 6 counties.")

wb.save(DST)
print("Saved.")
